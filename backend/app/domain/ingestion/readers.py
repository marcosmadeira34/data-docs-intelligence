"""Leitores tabulares em streaming para CSV e XLSX.

Nenhum leitor carrega o arquivo inteiro em memória: as linhas são produzidas
uma a uma. Números de linha são sempre 1-based, como o usuário vê no Excel.
"""

from __future__ import annotations

import codecs
import csv
import warnings
from collections import Counter
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.core.errors import FileRejected

csv.field_size_limit(1024 * 1024)

DELIMITER_CANDIDATES = [";", ",", "\t", "|"]
_CP1252_UNDEFINED = {0x81, 0x8D, 0x8F, 0x90, 0x9D}


@dataclass
class ReadOptions:
    sheet: str | None = None
    header_row: int | None = None
    delimiter: str | None = None
    encoding: str | None = None


@dataclass
class EncodingGuess:
    encoding: str
    confident: bool
    note: str | None = None


# --------------------------------------------------------------------------- CSV


def detect_encoding(path: Path) -> EncodingGuess:
    with path.open("rb") as fh:
        head = fh.read(4)
        if head.startswith(codecs.BOM_UTF8):
            return EncodingGuess("utf-8-sig", True)
        if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            return EncodingGuess("utf-16", True)
        fh.seek(0)
        decoder = codecs.getincrementaldecoder("utf-8")()
        try:
            while chunk := fh.read(1024 * 1024):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
            return EncodingGuess("utf-8", True)
        except UnicodeDecodeError:
            pass
        fh.seek(0)
        all_bytes = set()
        while chunk := fh.read(1024 * 1024):
            all_bytes.update(chunk)
    if all_bytes & _CP1252_UNDEFINED:
        return EncodingGuess("latin-1", False, "Codificação presumida ISO-8859-1 (latin-1); confirme os acentos na prévia.")
    return EncodingGuess("cp1252", False, "Codificação presumida Windows-1252; confirme os acentos na prévia.")


def _read_text_sample(path: Path, encoding: str, max_chars: int = 65536) -> str:
    with path.open("r", encoding=encoding, errors="replace", newline="") as fh:
        return fh.read(max_chars)


def score_delimiters(sample: str) -> list[dict[str, Any]]:
    """Avalia cada delimitador candidato pela consistência do número de campos."""
    lines = sample.splitlines()
    if len(lines) > 1 and len(sample) >= 65536:
        lines = lines[:-1]  # última linha pode estar truncada
    lines = [ln for ln in lines[:200] if ln.strip()]
    results = []
    for cand in DELIMITER_CANDIDATES:
        try:
            counts = [len(r) for r in csv.reader(lines, delimiter=cand)]
        except csv.Error:
            continue
        if not counts:
            continue
        modal, freq = Counter(counts).most_common(1)[0]
        if modal < 2:
            continue
        results.append({"delimiter": cand, "fields": modal, "consistency": round(freq / len(counts), 3)})
    # Consistência agrupada em décimos; dentro do grupo, mais campos vence
    # (ex.: ";" com 9 campos vence "," que só aparece na vírgula decimal).
    results.sort(key=lambda r: (round(r["consistency"], 1), r["fields"]), reverse=True)
    return results


def detect_delimiter(sample: str) -> tuple[str, list[dict[str, Any]], bool]:
    """Retorna (delimitador, candidatos, ambíguo)."""
    scored = score_delimiters(sample)
    if not scored:
        return ",", [], False  # arquivo de coluna única
    best = scored[0]
    ambiguous = False
    if len(scored) > 1:
        second = scored[1]
        ambiguous = (
            second["consistency"] >= 0.9
            and best["consistency"] - second["consistency"] < 0.05
            and second["fields"] >= best["fields"] * 0.75
        )
    return best["delimiter"], scored, ambiguous


@contextmanager
def _open_csv(path: Path, encoding: str, delimiter: str) -> Iterator[Iterator[list[str]]]:
    with path.open("r", encoding=encoding, errors="replace", newline="") as fh:
        yield csv.reader(fh, delimiter=delimiter)


def iter_csv_rows(path: Path, encoding: str, delimiter: str) -> Iterator[tuple[int, list[Any]]]:
    with _open_csv(path, encoding, delimiter) as reader:
        try:
            for idx, row in enumerate(reader, start=1):
                yield idx, row
        except csv.Error as exc:
            raise FileRejected("O CSV contém uma linha malformada que não pôde ser lida.", details={"line": reader.line_num}, code="malformed_csv") from exc


# --------------------------------------------------------------------------- XLSX


@contextmanager
def _workbook(path: Path, read_only: bool = True) -> Iterator[Any]:
    """Abre a planilha a partir de um file handle: chaves de armazenamento não têm extensão,
    e o openpyxl valida a extensão quando recebe um caminho."""
    import openpyxl

    with path.open("rb") as fh, warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = openpyxl.load_workbook(fh, read_only=read_only, data_only=True)
        except Exception as exc:  # openpyxl lança vários tipos para arquivos inválidos
            raise FileRejected("Não foi possível abrir a planilha .xlsx; o arquivo pode estar corrompido.", code="corrupted_file") from exc
        try:
            yield wb
        finally:
            wb.close()


def list_sheets(path: Path, *, inspect_merged: bool) -> list[dict[str, Any]]:
    with _workbook(path, read_only=True) as wb:
        sheets = []
        for ws in wb.worksheets:
            dims = ws.calculate_dimension(force=False) if hasattr(ws, "calculate_dimension") else None
            sheets.append({
                "name": ws.title,
                "dimension": dims if dims and dims != "A1:A1" else None,
                "visible": getattr(ws, "sheet_state", "visible") == "visible",
                "merged_ranges": None,
            })
    if inspect_merged:
        with _workbook(path, read_only=False) as full:
            for info, ws in zip(sheets, full.worksheets, strict=False):
                info["merged_ranges"] = [str(r) for r in list(ws.merged_cells.ranges)[:50]]
    return sheets


def iter_xlsx_rows(path: Path, sheet: str | None) -> Iterator[tuple[int, list[Any]]]:
    with _workbook(path, read_only=True) as wb:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise FileRejected("A aba informada não existe na planilha.", code="sheet_not_found", details={"sheets": wb.sheetnames})
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(row)
            while values and (values[-1] is None or values[-1] == ""):
                values.pop()
            yield idx, values


# --------------------------------------------------------------------------- API comum


@dataclass
class ResolvedOptions:
    format: str
    sheet: str | None
    encoding: str | None
    delimiter: str | None
    header_row: int


def iter_raw_rows(path: Path, fmt: str, opts: ResolvedOptions) -> Iterator[tuple[int, list[Any]]]:
    if fmt == "csv":
        assert opts.encoding and opts.delimiter
        return iter_csv_rows(path, opts.encoding, opts.delimiter)
    return iter_xlsx_rows(path, opts.sheet)


def build_columns(header_cells: list[Any]) -> tuple[list[str], list[dict[str, Any]]]:
    """Gera nomes de coluna únicos a partir do cabeçalho e relata problemas."""
    from app.domain.values import normalize_name, to_text

    names: list[str] = []
    issues: list[dict[str, Any]] = []
    seen: Counter[str] = Counter()
    for idx, cell in enumerate(header_cells):
        name = to_text(cell).replace("\n", " ").strip()
        if not name:
            name = f"Coluna {idx + 1}"
            issues.append({"code": "empty_header", "column": idx + 1, "message": f"A coluna {idx + 1} não tem nome no cabeçalho."})
        key = normalize_name(name)
        seen[key] += 1
        if seen[key] > 1:
            issues.append({"code": "duplicate_column", "column": idx + 1,
                           "message": f"Coluna duplicada: '{name}'. Renomeada para '{name} ({seen[key]})'."})
            name = f"{name} ({seen[key]})"
        names.append(name)
    return names, issues


def iter_records(path: Path, opts: ResolvedOptions) -> tuple[list[str], Iterator[tuple[int, list[Any]]]]:
    """Retorna (colunas, iterador de (linha, valores)) a partir da linha de cabeçalho."""
    rows = iter_raw_rows(path, opts.format, opts)
    header: list[Any] | None = None
    for idx, row in rows:
        if idx == opts.header_row:
            header = row
            break
    if header is None:
        raise FileRejected("A linha de cabeçalho informada não existe no arquivo.", code="header_not_found")
    columns, _ = build_columns(header)

    def gen() -> Iterator[tuple[int, list[Any]]]:
        for idx, row in rows:
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
                continue
            yield idx, row

    return columns, gen()
