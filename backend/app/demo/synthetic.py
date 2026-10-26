"""Gerador determinístico de arquivos sintéticos para demonstração e testes.

Todos os nomes, documentos e valores são fictícios. Os CNPJs são gerados com
dígitos verificadores válidos a partir de números aleatórios e não se referem
a empresas reais de forma intencional.

Uso: python -m app.demo.synthetic <diretório de saída>
"""

from __future__ import annotations

import csv
import random
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl

SEED = 20240901

SUPPLIERS = [
    "Alfa Suprimentos Ltda", "Beta Serviços de Limpeza ME", "Gama Tecnologia S.A.", "Delta Transportes Ltda",
    "Épsilon Papelaria EIRELI", "Zeta Consultoria Ltda", "Eta Energia S.A.", "Teta Alimentos Ltda",
    "Iota Manutenção Predial", "Capa Telecom S.A.", "Lambda Contabilidade Ltda", "Mi Segurança Patrimonial",
]
COST_CENTERS = ["ADM", "COM", "FIN", "OPS", "TI"]


def _cnpj(rng: random.Random) -> str:
    base = [rng.randint(0, 9) for _ in range(8)] + [0, 0, 0, 1]
    for weights in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        rest = sum(d * w for d, w in zip(base, weights, strict=False)) % 11
        base.append(0 if rest < 2 else 11 - rest)
    d = "".join(map(str, base))
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def _br(value: Decimal) -> str:
    s = f"{value:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _payables(period: date, count: int, seed: int) -> list[dict]:
    rng = random.Random(seed)
    cnpjs = {s: _cnpj(rng) for s in SUPPLIERS}
    rows = []
    for i in range(count):
        supplier = rng.choice(SUPPLIERS)
        issue = period + timedelta(days=rng.randint(0, 27))
        due = issue + timedelta(days=rng.choice([10, 15, 30]))
        value = Decimal(rng.randint(5000, 2500000)) / 100
        paid = rng.random() < 0.6
        rows.append({
            "doc": f"CP-{period:%Y%m}-{i + 1:04d}",
            "supplier": supplier,
            "cnpj": cnpjs[supplier],
            "issue": issue,
            "due": due,
            "payment": due - timedelta(days=rng.randint(0, 3)) if paid else None,
            "value": value,
            "status": "Pago" if paid else "Aberto",
            "cc": rng.choice(COST_CENTERS),
        })
    return rows


def _inject_problems(rows: list[dict]) -> None:
    """Introduz inconsistências conhecidas (documentadas em README da demo)."""
    rows[3]["cnpj"] = ""                      # CNPJ ausente
    rows[5]["cnpj"] = "12.345.678/0001-00"    # CNPJ com dígito verificador inválido
    rows[8]["value"] = "abc"                  # valor não numérico
    rows[11]["issue"] = "31/02/2024"          # data inexistente
    rows[14]["value"] = Decimal("-150.00")    # valor negativo
    rows[17]["status"] = "Pago"               # pago sem data de pagamento
    rows[17]["payment"] = None
    rows[20]["supplier"] = ""                 # fornecedor ausente
    rows.append(dict(rows[22]))               # documento duplicado


def _fmt_date(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return v.strftime("%d/%m/%Y")


def _fmt_value(v) -> str:
    return v if isinstance(v, str) else _br(v)


STANDARD_HEADER = ["Documento", "Fornecedor", "CNPJ", "Emissão", "Vencimento", "Data Pagamento", "Valor", "Status", "Centro de Custo"]
ALT_HEADER = ["Situação", "Nº Doc.", "Razão Social", "CPF/CNPJ", "Vlr. Título (R$)", "Dt. Emissão", "Data Venc.", "Dt. Baixa", "C. Custo"]


def _std_row(r: dict) -> list[str]:
    return [r["doc"], r["supplier"], r["cnpj"], _fmt_date(r["issue"]), _fmt_date(r["due"]), _fmt_date(r["payment"]),
            _fmt_value(r["value"]), r["status"], r["cc"]]


def _write_csv(path: Path, header: list[str], rows: list[list], *, delimiter: str = ";", encoding: str = "utf-8") -> Path:
    with path.open("w", encoding=encoding, newline="") as fh:
        w = csv.writer(fh, delimiter=delimiter)
        w.writerow(header)
        w.writerows(rows)
    return path


def generate_demo_files(out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}

    # 1. Layout padrão, setembro, com inconsistências injetadas
    sep = _payables(date(2024, 9, 1), 60, SEED)
    _inject_problems(sep)
    files["contas_pagar_2024_09.csv"] = _write_csv(
        out_dir / "contas_pagar_2024_09.csv", STANDARD_HEADER, [_std_row(r) for r in sep])

    # 2. Mesmo layout, outubro: compatível com o perfil aprovado
    octo = _payables(date(2024, 10, 1), 45, SEED + 1)
    octo[2]["value"] = Decimal("0")
    files["contas_pagar_2024_10.csv"] = _write_csv(
        out_dir / "contas_pagar_2024_10.csv", STANDARD_HEADER, [_std_row(r) for r in octo])

    # 3. Layout alterado: coluna Valor renomeada e CNPJ removida (incompatível)
    changed_header = ["Documento", "Fornecedor", "Emissão", "Vencimento", "Data Pagamento", "Valor Líquido", "Status", "Centro de Custo", "Observação"]
    changed_rows = [[r["doc"], r["supplier"], _fmt_date(r["issue"]), _fmt_date(r["due"]), _fmt_date(r["payment"]),
                     _fmt_value(r["value"]), r["status"], r["cc"], ""] for r in octo]
    files["contas_pagar_2024_11_layout_alterado.csv"] = _write_csv(
        out_dir / "contas_pagar_2024_11_layout_alterado.csv", changed_header, changed_rows)

    # 4. Layout alternativo em XLSX: cabeçalhos diferentes, outra ordem, linhas de título, duas abas
    wb = openpyxl.Workbook()
    ws_info = wb.active
    ws_info.title = "Instruções"
    ws_info["A1"] = "Relatório exportado do sistema financeiro (dados fictícios)"
    ws_info["A2"] = "Use a aba Lançamentos."
    ws = wb.create_sheet("Lançamentos")
    ws["A1"] = "CONTAS A PAGAR — SETEMBRO/2024"
    ws.merge_cells("A1:I1")
    ws["A2"] = "Empresa Demonstração Ltda (fictícia)"
    ws.append(ALT_HEADER)
    alt = _payables(date(2024, 9, 1), 30, SEED + 2)
    alt[4]["value"] = Decimal("1234567.89")
    for r in alt:
        ws.append([r["status"], r["doc"], r["supplier"], r["cnpj"], float(r["value"]),
                   datetime.combine(r["issue"], datetime.min.time()), datetime.combine(r["due"], datetime.min.time()),
                   datetime.combine(r["payment"], datetime.min.time()) if r["payment"] else None, r["cc"]])
    path = out_dir / "contas_pagar_alternativo_2024_09.xlsx"
    wb.save(path)
    files[path.name] = path

    # 5. Par de comparação: ERP (CSV) x banco (XLSX)
    rng = random.Random(SEED + 3)
    erp = _payables(date(2024, 9, 1), 40, SEED + 4)
    erp_rows = [[r["doc"], r["supplier"], _br(r["value"]), _fmt_date(r["due"])] for r in erp]
    files["erp_titulos_2024_09.csv"] = _write_csv(
        out_dir / "erp_titulos_2024_09.csv", ["Documento", "Fornecedor", "Valor", "Vencimento"], erp_rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extrato"
    ws.append(["Nr Documento", "Favorecido", "Valor Pago", "Data Pagamento"])
    bank = []
    for idx, r in enumerate(erp):
        if idx in (2, 9):  # 2 títulos do ERP sem pagamento no banco (somente em A)
            continue
        value = r["value"]
        if idx == 5:
            value += Decimal("37.50")   # divergência de valor
        if idx == 6:
            value += Decimal("0.004")   # diferença dentro da tolerância
        due = r["due"]
        if idx == 12:
            due += timedelta(days=2)    # divergência de data
        bank.append([r["doc"], r["supplier"].upper(), float(value), datetime.combine(due, datetime.min.time())])
    bank.append(["CP-202409-9001", "Pagamento sem título", 999.90, datetime(2024, 9, 20)])  # somente em B
    bank.append(["CP-202409-9002", "Tarifa bancária", 45.00, datetime(2024, 9, 30)])         # somente em B
    bank.append(list(bank[15]))                                                           # chave duplicada em B
    bank.append([None, "Lançamento sem documento", 10.0, datetime(2024, 9, 15)])          # sem chave
    rng.shuffle(bank)
    for row in bank:
        ws.append(row)
    path = out_dir / "banco_pagamentos_2024_09.xlsx"
    wb.save(path)
    files[path.name] = path

    # 6. CSV em Windows-1252 com vírgula como delimitador
    rows = [[r["doc"], r["supplier"], _fmt_value(r["value"]).replace(".", ""), _fmt_date(r["due"])] for r in octo[:10]]
    files["fornecedores_cp1252.csv"] = _write_csv(
        out_dir / "fornecedores_cp1252.csv", ["Código", "Razão Social", "Valor", "Vencimento"], rows,
        delimiter=",", encoding="cp1252")
    return files


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "demo_data")
    for name, p in generate_demo_files(target).items():
        print(f"{name}: {p.stat().st_size} bytes")
